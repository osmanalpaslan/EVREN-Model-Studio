"""Hesap kapama raporu ölçümü için SAHTE veri üretici, beklenen sonuç (kural motoru) ve çıktı denetçisi.

Gerçek müşteri verisi KULLANILMAZ: sütun yapısı kullanıcının dosyalarıyla aynıdır (2026-10-21 profil
başlıkları), değerler uydurmadır. Her SKILL.md kuralı için bilerek bir tuzak satırı vardır.
openpyxl gerekir (proje venv'inde yok; sistem Python'unda var):

    python tests/hesap_kapama/hk_veri.py uret   <klasör>
    python tests/hesap_kapama/hk_veri.py denetle <klasör> <çıktı.xlsx>   → JSON sonuç (stdout)
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

LOG_METNI = ("Müşteri ödeme aşamasında sistemden çıkmıştır. Faaliyet dışı kaydı nedeniyle tüm ürünler "
             "için otomatik olarak iş açılmıştır.")
API = {12055, 24153}
KULLANICILAR = [  # id, kod, ad, soyad, tam ad  (kod güncel soyadı göstermez: tuzak)
    (12055, "webservis", "Web", "Servis", "Web Servis"),
    (24153, "oto_hesapkapama", "oto_hesapkapama", "", "oto_hesapkapama"),
    (12826, "nefise.gungor", "Nefise", "Sökmen", "Nefise Sökmen"),
    (13001, "ahmet.yilmaz", "Ahmet", "Yılmaz", "Ahmet Yılmaz"),
    (13002, "melis.oztan", "Melis", "Güler", "Melis Güler"),
    (13003, "çağrı.öztürk", "Çağrı", "Öztürk", "Çağrı Öztürk"),
]
BILINMEYEN_ID = 19999                     # sözlükte yok → yalnız id yazılmalı ve raporda bildirilmeli
URUNLER = {"E-Fatura": 1, "E-Arşiv": 3, "E-İrsaliye": 6, "E-Defter": 14, "E-Envanter Defteri": 27}
ACIK_STATU = {"Yeni", "Devam Ediyor", "Ertelendi", "Teklif Verildi", "Demo", "Kazanıldı"}   # Kazanıldı istisnası
ISLER_BASLIK = ["Oluşturma Tarihi", "Ekleyen", "İşi Açan Kodu", "İş Sonucu", "Id", "Hesap Uyum", "Kaynak",
                "Kaynak2", "İşyeri Kodu", "Kategori1", "Müşteri Adı", "TC Kimlik No", "Vergi No", "İşi Yapacak",
                "İstek Versiyon Adı", "Konu", "İş No", "Üst İş No", "Durum", "Başlık", "Açıklama",
                "Bekleme Nedeni", "Öncelik", "Ürün", "Firsat Kodu", "İş Tipi", "Müşteri Kodu", "Sorumlu Kodu",
                "Kaybetme Nedeni", "Müşteri Kaybetme Nedeni", "Son Değiştiren"]
KISMI_OZET = ["İş ID", "İş Kodu", "Oluşturma Tarihi", "Müşteri Adı", "VKN/TCKN", "Ürün", "Fırsat Kodu",
              "Açılış Kaynağı", "Açılma Türü", "Açan Kullanıcı", "İş Sonucu", "İş Durumu", "Kapama Türü",
              "Kapatan Kullanıcı", "Kaybetme Nedeni", "Müşteri Kaybetme Nedeni", "Fırsat Durumu",
              "Fırsat Durum Kodu", "Sözleşme Durumu", "Sözleşme ID", "CRS'de Ürün Kapalı mı?",
              "CRS Kapanan Ürün", "CRS Kapanış Tarihi", "Kapanması Gereken", "Analiz Önceliği"]
SAYFALAR = ["Yönetim Özeti", "Kısmi Özet", "İş Listesi", "Sözleşme", "Fırsat", "CRS Kontrol", "Kapsam Dışı"]

# (sonuç, ekleyen, hesap_uyum, log, kaynak[tuzak], ürün, fırsat_statü, sözleşme, son_değiştiren,
#  kaybetme, müşteri_kaybetme, vkn_mi, crs_kapanis)
#  sözleşme: "id" / "#N/A" / "yok" / "cift" (iki farklı id = çoğullama) / "tekrar" (birebir aynı iki satır)
SENARYO = [
    # --- açık işler (16)
    ("Acik", 24153, True, True, "CRMIs", "E-Fatura", "Devam Ediyor", "id", 24153, "", "", True, False),     # Log
    ("Acik", 12055, True, True, "0", "E-Arşiv", "Yeni", "#N/A", 12055, "", "", True, False),                # Log
    ("Acik", 24153, True, False, "CRMTalep", "E-Fatura", "Kazanıldı", "id", 24153, "", "", True, False),     # HesapUyum
    ("Acik", 12826, True, False, "CRMIs", "E-Defter", "Teklif Verildi", "yok", 12826, "", "", False, False),  # HesapUyum+personel
    ("Acik", 13001, True, False, "CRMIs", "E-İrsaliye", "Fes Edildi", "id", 13001, "", "", True, False),     # HesapUyum+personel
    ("Acik", 24153, True, False, "", "E-Defter", "Ertelendi", "#N/A", 24153, "", "", True, False),
    ("Acik", 12055, False, False, "CRMTalep", "E-Fatura", "Demo", "id", 12055, "", "", True, False),        # Eski Destek
    ("Acik", 12055, False, False, "CRMTalep", "E-Arşiv", "Kaybedildi", "yok", 12055, "", "", False, False),  # Eski Destek
    ("Acik", 13002, False, False, "", "E-Fatura", "Yeni", "id", 13002, "", "", True, False),               # Manuel, Kaynak boş
    ("Acik", 13003, False, False, "CRMIs", "E-Envanter Defteri", "Devam Ediyor", "#N/A", 13003, "", "", True, False),
    ("Acik", BILINMEYEN_ID, False, False, "0", "E-İrsaliye", "Sehven Açıldı", "yok", BILINMEYEN_ID, "", "", True, False),
    ("Acik", 12826, False, False, "CRMTalep", "E-Fatura", "Potansiyel Değil", "cift", 12826, "", "", True, False),
    ("Acik", 24153, True, False, "CRMIs", "E-Arşiv", "Paket Değişimi", "id", 24153, "", "", True, True),     # CRS kapalı ama iş açık
    ("Acik", 13001, False, False, "CRMIs", "E-Defter", "Hatalı Fırsat", "yok", 13001, "", "", False, False),
    ("Acik", 24153, True, False, "CRMTalep", "E-Fatura", "Teklif Verildi", "tekrar", 24153, "", "", True, False),
    ("Acik", 13002, True, False, "", "E-İrsaliye", "Kaybedildi", "#N/A", 13002, "", "", True, False),
    # --- kapalı işler (16)
    ("Kapali", 24153, True, True, "CRMIs", "E-Fatura", "Kaybedildi", "#N/A", 24153, "FirmaFaaliyetDışı", "0", True, True),   # Log
    ("Kapali", 24153, True, False, "CRMTalep", "E-Fatura", "Kazanıldı", "id", 24153, "FirmaFaaliyetDışı", "0", True, True),   # 2 sorun
    ("Kapali", 24153, True, False, "CRMIs", "E-Arşiv", "Devam Ediyor", "#N/A", 24153, "GayriFaal", "0", True, True),        # 1 sorun (fırsat)
    ("Kapali", 12826, True, False, "CRMIs", "E-Defter", "Fes Edildi", "id", 12055, "Fiyat", "0", False, False),            # 1 sorun (sözleşme)
    ("Kapali", 12055, False, False, "CRMTalep", "E-İrsaliye", "Kaybedildi", "yok", 12055, "İhtiyaçKalmadı", "0", True, True),
    ("Kapali", 12055, False, False, "CRMTalep", "E-Envanter Defteri", "Vazgeçildi", "#N/A", 24153, "FirmaFaaliyetDışı", "0", True, True),  # CRS 127 (eşleme 0)
    ("Kapali", 13001, False, False, "", "E-Fatura", "Kaybedildi", "#N/A", 13001, "YönetimKararı", "0", True, False),
    ("Kapali", 13002, False, False, "CRMIs", "E-Arşiv", "Fırsat Yok", "id", 13003, "Fiyat", "ŞubeKapanışı", True, False),  # MÜŞTERİ kaybetme: kapsam İÇİ
    ("Kapali", 13003, False, False, "CRMIs", "E-Defter", "Teklif Verildi", "id", 13003, "DiğerNedenler", "0", False, False),  # 2 sorun
    ("Kapali", 24153, True, False, "CRMIs", "E-Fatura", "Mükellef Alındı", "#N/A", 24153, "PostaKutusuKapanışı", "0", True, True),  # kapsam dışı
    ("Kapali", 12826, False, False, "CRMTalep", "E-İrsaliye", "Kaybedildi", "yok", 13001, "ŞubeKapanışı", "0", True, False),       # kapsam dışı
    ("Kapali", 24153, True, False, "CRMIs", "E-Arşiv", "Kaybedildi", "#N/A", 24153, "FirmaFaaliyetDışı", "0", True, True),
    ("Kapali", 12055, False, False, "CRMTalep", "E-Fatura", "Birleştirildi", "#N/A", 12055, "GayriFaal", "0", True, False),
    ("Kapali", 13001, True, False, "CRMIs", "E-Defter", "Ertelendi", "#N/A", 24153, "FirmaFaaliyetDışı", "0", False, True),     # 1 sorun (fırsat)
    ("Kapali", 24153, True, False, "CRMTalep", "E-İrsaliye", "Kaybedildi", "#N/A", 24153, "GayriFaal", "0", True, True),
    ("Kapali", 13002, False, False, "0", "E-Fatura", "Fes Edildi", "#N/A", 13002, "Fiyat", "0", True, False),
]


def _kullanici(uid: int) -> str:
    for k in KULLANICILAR:
        if k[0] == uid:
            return f"{uid} - {k[4]}"
    return str(uid)


def kayitlar() -> list[dict]:
    base = dt.date(2026, 9, 1)
    out = []
    for i, s in enumerate(SENARYO):
        (sonuc, ekleyen, hu, log, kaynak, urun, statu, soz, son, kayb, mkayb, vkn_mi, crs) = s
        vkn = f"00{12345600 + i:08d}" if vkn_mi else ""          # baştaki sıfırlı 10 haneli VKN
        tckn = "" if vkn_mi else f"{30000000000 + i * 1111:011d}"
        out.append({
            "i": i, "sonuc": sonuc, "ekleyen": ekleyen, "hu": hu, "log": log, "kaynak": kaynak, "urun": urun,
            "statu": statu, "soz": soz, "son": son, "kayb": kayb, "mkayb": mkayb, "vkn": vkn, "tckn": tckn,
            "crs": crs, "id": 2100000 + i * 7, "is_no": f"DNS-2026{700000 + i}", "firsat": f"F-{900100 + i:09d}",
            "tarih": base + dt.timedelta(days=i), "musteri": f"Örnek Müşteri {i + 1} Ltd. Şti.",
            "soz_id": 41000000 + i * 13, "soz_id2": 41000000 + i * 13 + 5,
        })
    return out


def _vkn(r) -> str:
    return r["vkn"] or r["tckn"]


def beklenen(rows: list[dict]) -> dict:
    """SKILL.md kurallarının bağımsız uygulaması (kahin)."""
    acilis, acilma, kapama, ozet, kapsam_disi, crs_kontrol = {}, {}, {}, [], [], []
    for r in rows:
        if r["log"]:
            ak = "Log"
        elif r["hu"]:
            ak = "HesapUyum"
        elif r["ekleyen"] == 12055:
            ak = "Eski Destek Merkezi"
        else:
            ak = "Manuel"
        at = "Otomatik açılan iş" if r["ekleyen"] in API else "Manuel açılan iş"
        acilis[ak] = acilis.get(ak, 0) + 1
        acilma[at] = acilma.get(at, 0) + 1
        kapali = r["sonuc"] == "Kapali"
        if kapali:
            kt = {24153: "oto hesapkapama", 12055: "webservis"}.get(r["son"], "Manuel / Kullanıcı")
            kapama[kt] = kapama.get(kt, 0) + 1
        else:
            kt = "(iş açık)"
        f_acik = r["statu"] in ACIK_STATU
        soz_ids = {"id": [r["soz_id"]], "cift": [r["soz_id"], r["soz_id2"]], "tekrar": [r["soz_id"]]}.get(r["soz"], [None])
        disari = r["kayb"] in ("PostaKutusuKapanışı", "ŞubeKapanışı")
        for sid in soz_ids:
            aktif = sid is not None
            sorun = (int(f_acik) + int(aktif)) if kapali else 0
            row = {"İş ID": r["id"], "Açılış Kaynağı": ak, "Açılma Türü": at, "Açan Kullanıcı": _kullanici(r["ekleyen"]),
                   "Kapama Türü": kt, "Kapatan Kullanıcı": _kullanici(r["son"]) if kapali else "",
                   "Fırsat Durumu": "Açık" if f_acik else "Kapalı", "Fırsat Durum Kodu": r["statu"],
                   "Sözleşme Durumu": "Aktif" if aktif else "Pasif", "Sözleşme ID": str(sid) if aktif else "",
                   "CRS'de Ürün Kapalı mı?": "Evet" if r["crs"] else "Hayır",
                   "Analiz Önceliği": ({2: "1 - Yüksek", 1: "2 - Orta", 0: "3 - Sorun yok"}[sorun] if kapali else None),
                   "Kaybetme Nedeni": r["kayb"], "Müşteri Kaybetme Nedeni": r["mkayb"]}
            (kapsam_disi if disari else ozet).append(row)
        if r["crs"]:
            crs_kontrol.append(r["id"])
    return {"is_sayisi": len(rows), "acik": sum(r["sonuc"] == "Acik" for r in rows),
            "kapali": sum(r["sonuc"] == "Kapali" for r in rows), "acilis": acilis, "acilma": acilma,
            "kapama": kapama, "ozet": ozet, "kapsam_disi_is": sorted({x["İş ID"] for x in kapsam_disi}),
            "crs_kontrol_is": sorted(crs_kontrol),
            "kapanmasi_gereken": sum(1 for x in ozet if x["Analiz Önceliği"] in ("1 - Yüksek", "2 - Orta")),
            "bilinmeyen_kullanici": BILINMEYEN_ID}


# ---- dosya üretimi ---------------------------------------------------------------------------------

def uret(klasor: Path) -> dict:
    from openpyxl import Workbook

    klasor.mkdir(parents=True, exist_ok=True)
    rows = kayitlar()
    for sonuc, ad in (("Acik", "Açık işler.xlsx"), ("Kapali", "Kapalı işler.xlsx")):
        wb = Workbook()
        ws = wb.active
        ws.title = "UyumExcel"
        ws.append(ISLER_BASLIK)
        for r in rows:
            if r["sonuc"] != sonuc:
                continue
            acik = LOG_METNI if r["log"] else f"{r['urun']} hizmeti için hesap kapama talebi.\tMüşteri arandı."
            ws.append([r["tarih"], r["ekleyen"], "", r["sonuc"], r["id"], "True" if r["hu"] else "False",
                       r["kaynak"], "", "MRK", "Hesap Kapama", r["musteri"], r["tckn"], r["vkn"], "",
                       "", "Hesap kapama", r["is_no"], "", "Tamamlandı" if sonuc == "Kapali" else "Açık",
                       f"{r['urun']} kapama", acik, "", "Normal", r["urun"], r["firsat"], "Kapama",
                       f"M{1000 + r['i']}", "S01", r["kayb"] or "", r["mkayb"] or "", r["son"]])
        ws.cell(row=2, column=3, value="nefise.gungor")             # İşi Açan Kodu ≠ Ekleyen (tuzak)
        wb.save(klasor / ad)
    wb = Workbook()
    ws = wb.active
    ws.title = "UyumExcel"
    ws.append(["TCKN/VKN", "FıRSAT NO", "ÜRÜN KOD", "DURUM KOD", "SATıŞ TIPI"])
    for r in rows:
        ws.append([_vkn(r), r["firsat"], r["urun"], r["statu"], "ST"])
    wb.save(klasor / "Fırsat listesi 15.10.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sayfa1"
    ws.append(["TC Kimlik No", "Vergi No", "Ürün", "Firsat Kodu", "tcvkn", "uzunluk", "tırnak", "tcvknürün",
               "sözleşme ıd"])
    for r in rows:
        if r["soz"] == "yok":
            continue
        ids = {"id": [r["soz_id"]], "#N/A": ["#N/A"], "cift": [r["soz_id"], r["soz_id2"]],
               "tekrar": [r["soz_id"], r["soz_id"]]}[r["soz"]]
        for sid in ids:
            tv = _vkn(r)
            ws.append([r["tckn"], r["vkn"], r["urun"], r["firsat"], tv, len(tv), f"'{tv}'", f"{tv}{r['urun']}", sid])
    wb.save(klasor / "Sözleşme.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "UyumExcel"
    ws.append(["Firma Kodu", "Firma Adı", "Sayaç Tip Kodu", "Sayaç Tip Adı", "Pasif", "Ürün Grubu", "Satış Tip Kodu",
               "Crs Sayaç Tip Kodu", "Crs Gönderici Posta Kutusu", "Crs Alıcı Posta Kutusu", "CRS Sayaç Tipi",
               "CRS Kapanış Sayaç Tipi", "Icon"])
    for n, (urun, kod) in enumerate(URUNLER.items(), 1):
        kapanis = 0 if urun == "E-Envanter Defteri" else kod + 100      # eski veri: 0 (tuzak; +100 kuralı)
        ws.append([1, "Uyumsoft", f"ST{n:02d}", urun, "Hayır", "E-Dönüşüm", "ST", kod, "", "", kod, kapanis, ""])
    ws.append([1, "Uyumsoft", "ST99", "E-Yedekleme", "Hayır", "Diğer", "ST", 0, "", "", 0, 0, ""])
    wb.save(klasor / "Sayaç tipleri.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["IDENTIFY_TAX_NO", "TYPE", "HASACTIVEACCOUNT", "EXECUTION_DATE", "CO_ID", "CREATE_DATE", "STATU",
               "START_DATE"])
    for r in rows:
        tv = _vkn(r).lstrip("0")                                       # CRS'te baştaki sıfırlar yok
        kod = URUNLER[r["urun"]]
        ws.append([int(tv), kod, 1, r["tarih"], 2334, r["tarih"], 1000, r["tarih"]])          # açılış
        if r["crs"]:
            ws.append([int(tv), kod + 100, 0, r["tarih"] + dt.timedelta(days=20), 2334, r["tarih"], 1000, r["tarih"]])
    wb.save(klasor / "CRS AKTIVASYON DATA.xlsx")
    (klasor / "kullanici_listesi.tsv").write_text(
        "KULLANICI ID\tKULLANICI KODU\tADI\tSOYADI\tTAM AD\n"
        + "".join(f"{a}\t{b}\t{c}\t{d}\t{e}\n" for a, b, c, d, e in KULLANICILAR), encoding="utf-8")
    exp = beklenen(rows)
    (klasor.parent / f"{klasor.name}_beklenen.json").write_text(json.dumps(exp, ensure_ascii=False, indent=1),
                                                                 encoding="utf-8")
    return exp


def referans(cikti: Path, bozuk: bool = False) -> None:
    """Denetçiyi sınamak için (R-25) kahinden rapor yazar. bozuk=True → eski hatalı 'Kaynak' kuralı."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = kayitlar()
    exp = beklenen(rows)
    by = {r["id"]: r for r in rows}
    wb = Workbook()
    ws0 = wb.active
    ws0.title = SAYFALAR[0]
    ws0.append(["Hesap Kapama Raporu — Yönetim Özeti"])
    ws0.append(["Kaynak iş", exp["is_sayisi"], "Açık", exp["acik"], "Kapalı", exp["kapali"]])
    for k, v in exp["acilis"].items():
        ws0.append(["Açılış Kaynağı", k, v])
    ws0.append(["Kapanması gereken", exp["kapanmasi_gereken"]])

    def sayfa(ad, basl, satirlar):
        w = wb.create_sheet(ad)
        w.append(basl)
        for s in satirlar:
            w.append(s)
        for c in w[1]:
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.font = Font(bold=True, color="FFFFFF")
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 20
    ozet = []
    for e in exp["ozet"]:
        r = by[e["İş ID"]]
        ak = e["Açılış Kaynağı"]
        if bozuk:
            ak = {"CRMIs": "CRMIs", "CRMTalep": "CRMTalep"}.get(r["kaynak"], "Kaynak bilgisi yok")
        ozet.append([e["İş ID"], r["is_no"], r["tarih"].isoformat(), r["musteri"], _vkn(r), r["urun"], r["firsat"], ak,
                     e["Açılma Türü"], e["Açan Kullanıcı"], r["sonuc"], "Kapalı" if r["sonuc"] == "Kapali" else "Açık",
                     e["Kapama Türü"], e["Kapatan Kullanıcı"], e["Kaybetme Nedeni"], e["Müşteri Kaybetme Nedeni"],
                     e["Fırsat Durumu"], e["Fırsat Durum Kodu"], e["Sözleşme Durumu"], e["Sözleşme ID"],
                     e["CRS'de Ürün Kapalı mı?"], r["urun"] if r["crs"] else "", "", "", e["Analiz Önceliği"] or ""])
    sayfa(SAYFALAR[1], KISMI_OZET, ozet)
    sayfa(SAYFALAR[2], ISLER_BASLIK + ["Kaynak Dosya"],
          [[r["id"]] + [""] * 30 + ["Açık işler.xlsx" if r["sonuc"] == "Acik" else "Kapalı işler.xlsx"] for r in rows])
    sayfa(SAYFALAR[3], ["Fırsat Kodu", "Sözleşme ID"], [[r["firsat"], r["soz_id"]] for r in rows])
    sayfa(SAYFALAR[4], ["Fırsat Kodu", "Durum"], [[r["firsat"], r["statu"]] for r in rows])
    sayfa(SAYFALAR[5], ["İş ID", "Ürün"], [[i, by[i]["urun"]] for i in exp["crs_kontrol_is"]])
    sayfa(SAYFALAR[6], ["İş ID", "Kaybetme Nedeni"], [[i, by[i]["kayb"]] for i in exp["kapsam_disi_is"]])
    wb.save(cikti)


# ---- denetim ---------------------------------------------------------------------------------------

def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def denetle(klasor: Path, cikti: Path) -> dict:
    from openpyxl import load_workbook

    exp = beklenen(kayitlar())
    hatalar: list[str] = []
    gecen: list[str] = []

    def kontrol(ad, ok, ayrinti=""):
        (gecen if ok else hatalar).append(ad + ("" if ok or not ayrinti else f": {ayrinti}"))
    if not cikti.is_file():
        return {"dosya": False, "hatalar": ["çıktı Excel dosyası yok"], "gecen": [], "puan": "0"}
    wb = load_workbook(cikti, data_only=True)
    kontrol("7 sayfa, doğru sıra", wb.sheetnames == SAYFALAR, f"{wb.sheetnames}")
    ws = wb["Kısmi Özet"] if "Kısmi Özet" in wb.sheetnames else None
    if ws is None:
        hatalar.append("Kısmi Özet sayfası yok")
        return {"dosya": True, "hatalar": hatalar, "gecen": gecen}
    head = [_s(c.value) for c in ws[1]]
    head = head[:max((i for i, h in enumerate(head) if h), default=-1) + 1]
    kontrol("Kısmi Özet 25 sütun, sıra birebir", head == KISMI_OZET, f"{len(head)} sütun: {head}")
    kontrol("başlık dondurulmuş", ws.freeze_panes in ("A2", "B2"), str(ws.freeze_panes))
    kontrol("otomatik filtre açık", bool(ws.auto_filter and ws.auto_filter.ref), "")
    fill = ws.cell(row=1, column=1).fill
    rgb = (fill.fgColor.rgb or "") if fill and fill.fgColor is not None else ""
    kontrol("başlık koyu dolgulu", fill.fill_type == "solid" and isinstance(rgb, str) and rgb[-6:-4] < "60", rgb)
    idx = {h: i for i, h in enumerate(head)}
    rows = [[_s(c.value) for c in r] for r in ws.iter_rows(min_row=2) if any(c.value is not None for c in r)]
    kontrol("Kısmi Özet satır sayısı (çoğullama dahil, kapsam dışı hariç)", len(rows) == len(exp["ozet"]),
            f"{len(rows)} / beklenen {len(exp['ozet'])}")
    # SKILL çelişkisi: Kapama Türü değeri "Manuel / Kullanıcı" kural 5'teki " / " yasağına takılır → ayrık
    slash = sum(1 for r in rows for v in r if " / " in v and v.replace(" ", "") != "Manuel/Kullanıcı")
    kontrol("' / ' ile birleştirilmiş hücre yok", slash == 0, f"{slash} hücre")
    turkce = any("Ş" in v or "ş" in v or "ı" in v for r in rows for v in r)
    bozuk = sum(1 for r in rows for v in r if "Ã" in v or "Ä" in v or "�" in v)
    kontrol("Türkçe karakterler bozulmamış", turkce and bozuk == 0, f"bozuk hücre {bozuk}")
    by_id: dict[str, list[list[str]]] = {}
    for r in rows:
        by_id.setdefault(r[idx["İş ID"]] if "İş ID" in idx else "", []).append(r)
    alanlar = ["Açılış Kaynağı", "Açılma Türü", "Açan Kullanıcı", "Kapama Türü", "Kapatan Kullanıcı", "Fırsat Durumu",
               "Fırsat Durum Kodu", "Sözleşme Durumu", "Sözleşme ID", "CRS'de Ürün Kapalı mı?", "Analiz Önceliği"]
    yanlis: dict[str, list[str]] = {a: [] for a in alanlar}
    for e in exp["ozet"]:
        cand = by_id.get(str(e["İş ID"]), [])
        if not cand:
            yanlis.setdefault("eksik iş", []).append(str(e["İş ID"]))
            continue
        # çoğul satırlarda Sözleşme ID ile eşleştir
        got = next((c for c in cand if "Sözleşme ID" in idx and c[idx["Sözleşme ID"]] == e["Sözleşme ID"]), cand[0])
        for a in alanlar:
            if a not in idx:
                continue
            want = e[a]
            if want is None:                       # açık işte Analiz Önceliği tanımsız: denetlenmez
                continue
            have = got[idx[a]]
            if a in ("Açan Kullanıcı", "Kapatan Kullanıcı"):
                ok = have.replace("–", "-") == want
            elif a == "Sözleşme ID":
                ok = (have in ("", "-", "—") and want == "") or have == want
            elif a == "Kapama Türü" and want != "(iş açık)":
                ok = have.replace(" ", "").casefold() == want.replace(" ", "").casefold()
            elif a == "Kapama Türü" and want == "(iş açık)":
                ok = "açık" in have.lower()
            else:
                ok = have.casefold() == want.casefold()
            if not ok:
                yanlis[a].append(f"{e['İş ID']}: '{have}' ≠ '{want}'")
    for a, lst in yanlis.items():
        kontrol(f"alan: {a}", not lst, "; ".join(lst[:4]) + (f" (+{len(lst) - 4})" if len(lst) > 4 else ""))
    if "Kapsam Dışı" in wb.sheetnames:
        kd = wb["Kapsam Dışı"]
        kd_rows = [r for r in kd.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
        kontrol("Kapsam Dışı: yalnız Kaybetme Nedeni kuralı (2 iş)", len(kd_rows) == len(exp["kapsam_disi_is"]),
                f"{len(kd_rows)} satır / beklenen {len(exp['kapsam_disi_is'])}")
    if "CRS Kontrol" in wb.sheetnames:
        cr = wb["CRS Kontrol"]
        cr_rows = [r for r in cr.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
        kontrol("CRS Kontrol satır sayısı", len(cr_rows) == len(exp["crs_kontrol_is"]),
                f"{len(cr_rows)} / beklenen {len(exp['crs_kontrol_is'])}")
    if "İş Listesi" in wb.sheetnames:
        il = wb["İş Listesi"]
        il_rows = [r for r in il.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
        kontrol("İş Listesi = açık + kapalı", len(il_rows) == exp["is_sayisi"], f"{len(il_rows)}")
        ih = [_s(c.value) for c in il[1]]
        kontrol("İş Listesi'nde 'Kaynak Dosya' sütunu", "Kaynak Dosya" in ih, "")
    for name in SAYFALAR:
        if name in wb.sheetnames:
            w = wb[name]
            genis = [round(d.width, 1) for d in w.column_dimensions.values() if d.width and d.width > 42.5]
            if genis:
                hatalar.append(f"{name}: sütun genişliği > 42 ({genis[:3]})")
    yo = wb["Yönetim Özeti"] if "Yönetim Özeti" in wb.sheetnames else None
    if yo is not None:
        metin = " ".join(_s(c) for r in yo.iter_rows(values_only=True) for c in r)
        kontrol("Yönetim Özeti dolu", len(metin) > 100, f"{len(metin)} karakter")
        kontrol("Yönetim Özeti'nde NOTLAR bloğu yok", "NOTLAR" not in metin.upper(), "")
    toplam = len(gecen) + len(hatalar)
    return {"dosya": True, "gecen": gecen, "hatalar": hatalar, "puan": f"{len(gecen)}/{toplam}"}


if __name__ == "__main__":
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    cmd = sys.argv[1]
    if cmd == "uret":
        e = uret(Path(sys.argv[2]))
        print(json.dumps({"is": e["is_sayisi"], "acilis": e["acilis"], "acilma": e["acilma"], "kapama": e["kapama"],
                          "ozet_satir": len(e["ozet"]), "kapsam_disi": e["kapsam_disi_is"],
                          "kapanmasi_gereken": e["kapanmasi_gereken"]}, ensure_ascii=False))
    elif cmd == "referans":                         # referans <çıktı.xlsx> [bozuk]
        referans(Path(sys.argv[2]), bozuk=len(sys.argv) > 3 and sys.argv[3] == "bozuk")
    elif cmd == "denetle":
        print(json.dumps(denetle(Path(sys.argv[2]), Path(sys.argv[3])), ensure_ascii=False, indent=1))
